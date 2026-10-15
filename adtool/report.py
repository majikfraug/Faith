"""Progress reports and the recording script."""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Sequence

from .config import Config
from .cues import Cue, CueSheet
from .timecode import fmt_short


def status_markdown(cfg: Config, sheet: CueSheet, block_s: float = 300.0) -> str:
    s = sheet.summary()
    L: List[str] = [f"# Status: {cfg.get('project_name', default='Film')}", ""]
    L.append(f"Cues: {s['total']}  (" + ", ".join(f"{k} {v}" for k, v in sorted(s['by_type'].items())) + ")")
    L.append("Status: " + ", ".join(f"{k} {v}" for k, v in sorted(s['by_status'].items())))
    L.append("Fit: " + ", ".join(f"{k} {v}" for k, v in sorted(s['by_fits'].items())))
    L.append(f"Description time available: {s['desc_seconds']:.0f}s in DESC windows; "
             f"subtitles estimated too long for their window: {s['sub_est_no']}")
    L.append("")
    L.append("| block | cues | new | drafted | approved | recorded | skip | fit=no |")
    L.append("|---|---|---|---|---|---|---|---|")
    end = sheet.duration_hint()
    next_range: Optional[str] = None
    b = 0.0
    while b < end:
        blk = [c for c in sheet.cues if b <= c.start_s < b + block_s]
        if blk:
            counts = {k: sum(1 for c in blk if c.status == k) for k in ("new", "drafted", "approved", "recorded", "skip")}
            no = sum(1 for c in blk if c.fits == "no" and c.status != "skip")
            L.append(f"| {fmt_short(b)}–{fmt_short(min(b + block_s, end))} | {len(blk)} | {counts['new']} | {counts['drafted']} | "
                     f"{counts['approved']} | {counts['recorded']} | {counts['skip']} | {no} |")
            if next_range is None and counts["new"] > 0:
                next_range = f"{fmt_short(b)}-{fmt_short(min(b + block_s, end))}"
        b += block_s
    L.append("")
    L.append(f"Next range to draft: {next_range}" if next_range else "All blocks have been drafted.")
    return "\n".join(L) + "\n"


TYPE_WORD = {"SUB": "Subtitle", "TEXT": "Text", "DESC": "Description"}


def script_cues(cues: Sequence[Cue], include_new: bool = False) -> List[Cue]:
    """The cues a reading script shows: every line with words, except skipped ones and (by default) undrafted ones."""
    keep = []
    for c in cues:
        if c.status == "skip" or not c.line.strip():
            continue
        if c.status == "new" and not include_new:
            continue
        keep.append(c)
    return keep


def script_text(cfg: Config, sheet: CueSheet, cues: Optional[Sequence[Cue]] = None, include_new: bool = False) -> str:
    """A plain-text reading script (Notepad-friendly): one line per cue, grouped by minute."""
    cues = script_cues(cues if cues is not None else sheet.cues, include_new)
    title = f"{cfg.get('project_name', default='Film')} - audio description script"
    L: List[str] = [title, "=" * len(title), "",
                    "Time      Type         Line   (* = recorded, ? = not yet drafted)", ""]
    current_min = -1
    for c in cues:
        m = int(c.start_s // 60)
        if m != current_min:
            current_min = m
            L.append(f"--- {m // 60:d}:{m % 60:02d}:00 ---" if m >= 60 else f"--- {m:02d}:00 ---")
        mark = "*" if c.status == "recorded" else ("?" if c.status == "new" else " ")
        L.append(f"{fmt_short(c.start_s):>9} {mark} {TYPE_WORD[c.type]:<11}  {c.line}")
    L.append("")
    return "\n".join(L) + "\n"


def script_docx(cfg: Config, sheet: CueSheet, out: Path, cues: Optional[Sequence[Cue]] = None,
                include_new: bool = False) -> Path:
    """A Word reading script: a heading per minute, then timecode, kind and line for each cue."""
    try:
        import docx  # python-docx
        from docx.shared import Pt, RGBColor
    except ImportError as e:  # pragma: no cover
        raise SystemExit("python-docx is not installed: .venv\\Scripts\\python -m pip install python-docx") from e
    cues = script_cues(cues if cues is not None else sheet.cues, include_new)
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)
    d.add_heading(f"{cfg.get('project_name', default='Film')} — audio description script", level=0)
    p = d.add_paragraph()
    p.add_run("Subtitle").bold = True
    p.add_run(" = read over the original speaker.  ")
    p.add_run("Text").bold = True
    p.add_run(" = on-screen text, read as written.  ")
    p.add_run("Description").bold = True
    p.add_run(" = what is on screen, in a gap.  A ✓ marks a line already recorded.")
    grey = RGBColor(0x60, 0x60, 0x60)
    current_min = -1
    for c in cues:
        m = int(c.start_s // 60)
        if m != current_min:
            current_min = m
            d.add_heading(f"{m // 60:d}:{m % 60:02d}:00" if m >= 60 else f"{m:02d}:00", level=2)
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        tag = p.add_run(f"{fmt_short(c.start_s)}  {TYPE_WORD[c.type]}")
        tag.bold = True
        tag.font.color.rgb = grey
        tag.font.size = Pt(10)
        if c.status == "recorded":
            p.add_run("  ✓").font.color.rgb = grey
        elif c.status == "new":
            p.add_run("  (not yet drafted)").font.color.rgb = grey
        d.add_paragraph(c.line).paragraph_format.left_indent = Pt(18)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out


def recording_script(cfg: Config, sheet: CueSheet, cues: Optional[Sequence[Cue]] = None) -> str:
    cues = [c for c in (cues if cues is not None else sheet.cues) if c.status != "skip" and c.line.strip()]
    L: List[str] = [f"# Recording script: {cfg.get('project_name', default='Film')}", ""]
    L.append("How to record in one pass: read each line, then stay silent for at least two seconds before the next one. "
             "Do not read the cue numbers aloud. If you flub a line, pause two seconds and read it again; "
             "then tell Claude which cue to keep. Keep the pace natural; the 'max' time is the window you must fit.")
    L.append("")
    glossary = cfg.project / "glossary.md"
    if glossary.exists():
        L.append("## Names and pronunciations (from project/glossary.md)\n")
        L.append(glossary.read_text(encoding="utf-8").strip())
        L.append("")
    L.append("## Lines\n")
    for c in cues:
        tag = {"SUB": "subtitle", "TEXT": "on-screen text", "DESC": "description"}[c.type]
        L.append(f"{c.id:04d}  {fmt_short(c.start_s)}  max {c.window_s:.1f}s  ({tag})")
        L.append(f"      {c.line}")
        L.append("")
    return "\n".join(L) + "\n"
